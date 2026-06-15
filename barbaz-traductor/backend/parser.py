import fitz  # PyMuPDF
from docx import Document
import io


def parse_file(file_bytes: bytes, filename: str) -> str:
    """Parse PDF, DOCX, or TXT file bytes and return plain text."""
    filename_lower = filename.lower()

    if filename_lower.endswith(".pdf"):
        return _parse_pdf(file_bytes)
    elif filename_lower.endswith(".docx"):
        return _parse_docx(file_bytes)
    elif filename_lower.endswith(".txt"):
        return _parse_txt(file_bytes)
    else:
        raise ValueError(f"Unsupported file type: {filename}. Supported types: PDF, DOCX, TXT")


def _parse_pdf(file_bytes: bytes) -> str:
    """Extract text from PDF using PyMuPDF."""
    text_parts = []
    with fitz.open(stream=file_bytes, filetype="pdf") as doc:
        for page_num, page in enumerate(doc, start=1):
            page_text = page.get_text("text")
            if page_text.strip():
                text_parts.append(f"[Página {page_num}]\n{page_text}")

    full_text = "\n\n".join(text_parts)
    if not full_text.strip():
        raise ValueError("No se pudo extraer texto del PDF. El documento puede estar escaneado o protegido.")
    return full_text


def _parse_docx(file_bytes: bytes) -> str:
    """Extract text from DOCX using python-docx."""
    doc = Document(io.BytesIO(file_bytes))
    text_parts = []

    for paragraph in doc.paragraphs:
        if paragraph.text.strip():
            text_parts.append(paragraph.text)

    for table in doc.tables:
        for row in table.rows:
            row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
            if row_text:
                text_parts.append(row_text)

    full_text = "\n\n".join(text_parts)
    if not full_text.strip():
        raise ValueError("No se pudo extraer texto del documento DOCX.")
    return full_text


def _parse_txt(file_bytes: bytes) -> str:
    """Decode TXT file bytes to string."""
    encodings = ["utf-8", "latin-1", "cp1252", "iso-8859-1"]
    for encoding in encodings:
        try:
            text = file_bytes.decode(encoding)
            if text.strip():
                return text
        except (UnicodeDecodeError, ValueError):
            continue
    raise ValueError("No se pudo decodificar el archivo de texto. Verifique la codificación del archivo.")
